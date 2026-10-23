"""Independent text-span recall, not a claim of complete layout accuracy."""
import datetime as dt
import json
from pathlib import Path
import re
import statistics
import sys
import time
import unicodedata
import zipfile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation


def norm(s):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(s)))


def references(item):
    path, fmt = item['path'], item['format']
    values = []
    if fmt == 'docx':
        doc = Document(path)
        def blocks(container):
            for block in container.iter_inner_content():
                if isinstance(block, Paragraph):
                    values.append(block.text)
                elif isinstance(block, Table):
                    seen = set()
                    for row in block.rows:
                        for cell in row.cells:
                            if cell._tc in seen:
                                continue
                            seen.add(cell._tc)
                            blocks(cell)
        blocks(doc)
        for rel in doc.part.rels.values():
            if rel.reltype.endswith(('footnotes', 'endnotes')) and not rel.is_external:
                tree = etree.fromstring(rel.target_part.blob)
                for p in tree.xpath('//*[local-name()="p"]'):
                    values.append(''.join(p.xpath('.//*[local-name()="t"]/text()')))
        source = 'python-docx paragraphs/tables + lxml footnotes'
    elif fmt == 'xlsx':
        from openpyxl.styles.numbers import is_date_format
        book = load_workbook(path, data_only=True)
        for sheet in book:
            for row in sheet:
                for cell in row:
                    v = cell.value
                    if v is None:
                        continue
                    if isinstance(v, dt.datetime):
                        v = v.strftime('%Y-%m-%d %H:%M:%S' if v.time() != dt.time() else '%Y-%m-%d')
                    elif isinstance(v, dt.date):
                        v = v.isoformat()
                    elif isinstance(v, dt.time):
                        # openpyxl maps zero serial to time even for a date format.
                        # Restore the workbook epoch for the 1904 zero-date case.
                        code=re.sub(r'"[^"]*"|\\.', '',cell.number_format.lower())
                        if book.epoch.year==1904 and is_date_format(cell.number_format) and re.search('[yd]',code):
                            v=book.epoch.strftime('%Y-%m-%d')
                        else:
                            v = v.strftime('%H:%M:%S')
                    elif isinstance(v, dt.timedelta):
                        seconds = round(v.total_seconds())
                        v = f'{seconds // 3600}:{seconds // 60 % 60:02}:{seconds % 60:02}'
                    elif isinstance(v, bool):
                        v = 'TRUE' if v else 'FALSE'
                    values.append(str(v))
        source = 'openpyxl cached values; dates normalized ISO; no formula evaluation'
    elif fmt == 'pptx':
        pres = Presentation(path)
        def shapes(items):
            for shape in items:
                if shape.has_text_frame:
                    values.extend(p.text for p in shape.text_frame.paragraphs)
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            values.extend(p.text for p in cell.text_frame.paragraphs)
                if hasattr(shape, 'shapes'):
                    shapes(shape.shapes)
        for slide in pres.slides:
            shapes(slide.shapes)
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame:
                    values.extend(p.text for p in frame.paragraphs)
        source = 'python-pptx text frames/tables/notes'
    else:
        with zipfile.ZipFile(path) as z:
            for part in z.namelist():
                if re.fullmatch(r'Contents/section\d+\.xml', part):
                    tree = etree.fromstring(z.read(part))
                    # Direct text runs are independently identified from the XML.
                    # This scores text recovery, not table order, layout or images.
                    for p in tree.xpath('//*[local-name()="p"]'):
                        values.append(''.join(p.xpath('./*[local-name()="run"]/*[local-name()="t"]//text()')))
        source = 'independent lxml HWPX direct paragraph runs; text-only coverage'
    return list(dict.fromkeys(norm(v) for v in values if norm(v))), source

